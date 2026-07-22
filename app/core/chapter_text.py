"""Helpers for turning an uploaded chapter file into plain, editable text.

This is the piece that was missing before: chapter uploads (particularly
.docx files) were never actually parsed into text, so authors ended up with
"chapters" that had no editable content. Every chapter-producing code path
(single upload, bulk upload, admin publish assembly) should go through
`extract_chapter_text` / `plain_text_to_html_paragraphs` below so behaviour
stays consistent.
"""

from __future__ import annotations

import io
import re

import bleach

CHAPTER_ALLOWED_EXTENSIONS = {".txt", ".docx"}
CHAPTER_FILENAME_FORMAT = (
    "Name chapter files like 'Chapter 01.docx', 'Chapter 01 - Arrival.docx', "
    "or '01_Arrival.txt'. Leading zeros are optional."
)
CHAPTER_FILENAME_PATTERN = re.compile(
    r"""
    ^\s*
    (?:
        (?:chapter|chap|ch)\s*
    )?
    (?P<number>0*[1-9]\d*)
    (?:
        \s*(?:[-_:.\u2013\u2014])\s*
        (?P<title>.+?)
    )?
    \s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)


class ChapterExtractionError(ValueError):
    """Raised when a chapter file can't be turned into usable text."""


def normalize_text(value: str) -> str:
    text = (value or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\u00a0\u200b\ufeff]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_chapter_filename(filename: str) -> tuple[int, str | None]:
    """Return the chapter number and optional title encoded in a filename.

    Supported examples:
    - Chapter 01.docx
    - Chapter 1 - The Door Opens.txt
    - Ch 02_Second Night.docx
    - 03.Arrival.txt
    """
    stem = filename.rsplit(".", 1)[0] if "." in filename else filename
    match = CHAPTER_FILENAME_PATTERN.match(stem)
    if not match:
        raise ChapterExtractionError(
            f"'{filename}' does not match the required chapter filename format. "
            f"{CHAPTER_FILENAME_FORMAT}"
        )

    title = normalize_text(re.sub(r"[_\-]+", " ", match.group("title") or ""))
    return int(match.group("number")), (title[:255] if title else None)


def extract_docx_text(content: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency should always be installed
        raise ChapterExtractionError(
            "DOCX support is not available on the server right now."
        ) from exc

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise ChapterExtractionError(
            "That .docx file looks corrupted or isn't a real Word document."
        ) from exc

    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    # Word documents sometimes carry all their text inside tables instead of
    # (or in addition to) top-level paragraphs.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    text = "\n\n".join(normalize_text(part) for part in paragraphs if normalize_text(part))
    return normalize_text(text)


def extract_txt_text(content: bytes) -> str:
    return normalize_text(content.decode("utf-8", errors="replace"))


def extract_chapter_text(filename: str, extension: str, content: bytes) -> str:
    """Extract plain text from an uploaded chapter file.

    Raises `ChapterExtractionError` with a user-facing message on failure.
    """
    if extension == ".docx":
        text = extract_docx_text(content)
    elif extension == ".txt":
        text = extract_txt_text(content)
    else:
        raise ChapterExtractionError(f"Unsupported file type '{extension}'.")

    if not text:
        raise ChapterExtractionError(
            f"'{filename}' doesn't contain any readable text."
        )

    return text


def chapter_title_from_filename(filename: str) -> str:
    try:
        number, parsed_title = parse_chapter_filename(filename)
        return (parsed_title or f"Chapter {number}")[:255]
    except ChapterExtractionError:
        pass

    stem = filename.rsplit(".", 1)[0] if "." in filename else filename
    stem = re.sub(r"[_\-]+", " ", stem).strip()
    stem = re.sub(r"\s+", " ", stem)
    return stem.title()[:255] if stem else "Untitled Chapter"


def word_count(text: str) -> int:
    return len((text or "").split())


def plain_text_to_html_paragraphs(text: str) -> str:
    """Convert plain chapter text into sanitized `<p>` paragraphs for reading."""
    paragraphs = [
        normalize_text(part)
        for part in re.split(r"\n{2,}", text or "")
        if normalize_text(part)
    ]
    return "".join(
        f"<p>{bleach.clean(part, tags=[], strip=True)}</p>" for part in paragraphs
    )
